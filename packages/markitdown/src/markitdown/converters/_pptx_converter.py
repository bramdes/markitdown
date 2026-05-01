import sys
import os
import io
import re
import html
import tempfile

from typing import BinaryIO, Any

from ._html_converter import HtmlConverter
from ._llm_caption import llm_caption
from .._base_converter import DocumentConverter, DocumentConverterResult
from .._stream_info import StreamInfo
from .._exceptions import MissingDependencyException, MISSING_DEPENDENCY_MESSAGE

# Try loading optional (but in this case, required) dependencies
# Save reporting of any exceptions for later
_dependency_exc_info = None
try:
    import pptx
except ImportError:
    # Preserve the error and stack trace for later
    _dependency_exc_info = sys.exc_info()


ACCEPTED_MIME_TYPE_PREFIXES = [
    "application/vnd.openxmlformats-officedocument.presentationml",
]

ACCEPTED_FILE_EXTENSIONS = [".pptx"]


class PptxConverter(DocumentConverter):
    """
    Converts PPTX files to Markdown.

    Text frames, tables, charts (with structured data), and pictures are
    extracted shape-by-shape. Pictures are transcribed via the vision LLM
    (when configured) and the result is inlined as text — no image references
    are emitted. Slides that contain non-text visual shapes (AutoShapes,
    SmartArt, FreeForm, ink, non-trivial groups) are additionally rendered
    to PNG via PowerPoint COM and transcribed by the same LLM, with the
    result appended to the slide's markdown.
    """

    def __init__(self):
        super().__init__()
        self._html_converter = HtmlConverter()

    def accepts(
        self,
        file_stream: BinaryIO,
        stream_info: StreamInfo,
        **kwargs: Any,
    ) -> bool:
        mimetype = (stream_info.mimetype or "").lower()
        extension = (stream_info.extension or "").lower()

        if extension in ACCEPTED_FILE_EXTENSIONS:
            return True

        for prefix in ACCEPTED_MIME_TYPE_PREFIXES:
            if mimetype.startswith(prefix):
                return True

        return False

    def convert(
        self,
        file_stream: BinaryIO,
        stream_info: StreamInfo,
        **kwargs: Any,
    ) -> DocumentConverterResult:
        if _dependency_exc_info is not None:
            raise MissingDependencyException(
                MISSING_DEPENDENCY_MESSAGE.format(
                    converter=type(self).__name__,
                    extension=".pptx",
                    feature="pptx",
                )
            ) from _dependency_exc_info[
                1
            ].with_traceback(  # type: ignore[union-attr]
                _dependency_exc_info[2]
            )

        llm_client = kwargs.get("llm_client")
        llm_model = kwargs.get("llm_model")
        llm_prompt = kwargs.get("llm_prompt")
        vision_enabled = llm_client is not None and llm_model is not None

        # Read the full stream once so we can both parse it and (if needed)
        # write it back out to disk for PowerPoint COM to open.
        pptx_bytes = file_stream.read()
        presentation = pptx.Presentation(io.BytesIO(pptx_bytes))

        slide_sections: list[str] = []
        slides_needing_render: list[int] = []

        for slide_index, slide in enumerate(presentation.slides, start=1):
            section_parts: list[str] = [f"\n\n<!-- Slide number: {slide_index} -->\n"]
            needs_render_flag = {"value": False}
            title = slide.shapes.title

            def get_shape_content(shape):
                if self._is_picture(shape):
                    section_parts.append(
                        self._render_picture(
                            shape,
                            llm_client=llm_client,
                            llm_model=llm_model,
                            llm_prompt=llm_prompt,
                        )
                    )
                    return

                if self._is_table(shape):
                    section_parts.append(
                        self._convert_table_to_markdown(shape.table, **kwargs)
                    )
                    return

                if shape.has_chart:
                    chart_md, chart_ok = self._convert_chart_to_markdown(shape.chart)
                    section_parts.append(chart_md)
                    if not chart_ok:
                        needs_render_flag["value"] = True
                    return

                if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
                    sorted_subshapes = sorted(
                        shape.shapes,
                        key=lambda x: (
                            float("-inf") if not x.top else x.top,
                            float("-inf") if not x.left else x.left,
                        ),
                    )
                    for subshape in sorted_subshapes:
                        get_shape_content(subshape)
                    return

                if self._is_visual_shape(shape):
                    needs_render_flag["value"] = True

                if shape.has_text_frame:
                    if shape == title:
                        section_parts.append("# " + shape.text.lstrip() + "\n")
                    else:
                        section_parts.append(shape.text + "\n")

            sorted_shapes = sorted(
                slide.shapes,
                key=lambda x: (
                    float("-inf") if not x.top else x.top,
                    float("-inf") if not x.left else x.left,
                ),
            )
            for shape in sorted_shapes:
                get_shape_content(shape)

            if slide.has_notes_slide:
                section_parts.append("\n\n### Notes:\n")
                notes_frame = slide.notes_slide.notes_text_frame
                if notes_frame is not None:
                    section_parts.append(notes_frame.text)

            slide_sections.append("".join(section_parts).strip())
            if needs_render_flag["value"] and vision_enabled:
                slides_needing_render.append(slide_index)

        if slides_needing_render:
            transcriptions = self._transcribe_slides(
                pptx_bytes,
                slides_needing_render,
                client=llm_client,
                model=llm_model,
                prompt=llm_prompt,
            )
            for idx, text in transcriptions.items():
                if not text:
                    continue
                pos = idx - 1
                if 0 <= pos < len(slide_sections):
                    slide_sections[pos] = (
                        slide_sections[pos].rstrip()
                        + "\n\n"
                        + text.strip()
                    )

        return DocumentConverterResult(
            markdown="\n".join(slide_sections).strip()
        )

    def _render_picture(
        self,
        shape,
        *,
        llm_client,
        llm_model,
        llm_prompt,
    ) -> str:
        alt_text = ""
        try:
            alt_text = shape._element._nvXxPr.cNvPr.attrib.get("descr", "") or ""
        except Exception:
            pass
        alt_text = re.sub(r"[\r\n\[\]]", " ", alt_text)
        alt_text = re.sub(r"\s+", " ", alt_text).strip()

        if llm_client is not None and llm_model is not None:
            image_filename = shape.image.filename
            image_extension = (
                os.path.splitext(image_filename)[1] if image_filename else None
            )
            image_stream_info = StreamInfo(
                mimetype=shape.image.content_type,
                extension=image_extension,
                filename=image_filename,
            )
            try:
                description = llm_caption(
                    io.BytesIO(shape.image.blob),
                    image_stream_info,
                    client=llm_client,
                    model=llm_model,
                    prompt=llm_prompt,
                )
            except Exception:
                description = None
            if description and description.strip():
                return "\n" + description.strip() + "\n"

        label = alt_text or shape.name or "image"
        return f"\n[Image: {label}]\n"

    def _is_picture(self, shape):
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            return True
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PLACEHOLDER:
            if hasattr(shape, "image"):
                return True
        return False

    def _is_table(self, shape):
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.TABLE:
            return True
        return False

    @staticmethod
    def _is_visual_shape(shape) -> bool:
        """Return True for shapes whose visual meaning isn't captured by the
        text/table/chart/picture extractors. A slide containing any of these
        is a candidate for full-slide LLM transcription."""
        try:
            t = shape.shape_type
        except Exception:
            return False
        MSO = pptx.enum.shapes.MSO_SHAPE_TYPE
        candidates = {
            getattr(MSO, name, None)
            for name in ("AUTO_SHAPE", "FREEFORM", "DIAGRAM", "INK")
        }
        candidates.discard(None)
        return t in candidates

    def _convert_table_to_markdown(self, table, **kwargs):
        # Write the table as HTML, then convert it to Markdown
        html_table = "<html><body><table>"
        first_row = True
        for row in table.rows:
            html_table += "<tr>"
            for cell in row.cells:
                if first_row:
                    html_table += "<th>" + html.escape(cell.text) + "</th>"
                else:
                    html_table += "<td>" + html.escape(cell.text) + "</td>"
            html_table += "</tr>"
            first_row = False
        html_table += "</table></body></html>"

        markdown_result = self._html_converter.convert_string(html_table, **kwargs).markdown.strip()
        markdown_result = self._clean_empty_table_headers(markdown_result)
        return markdown_result + "\n"

    def _clean_empty_table_headers(self, markdown_table):
        """Remove empty table header rows and their separators."""
        lines = markdown_table.split('\n')
        cleaned_lines = []
        i = 0

        while i < len(lines):
            line = lines[i].strip()

            if line and all(c in '| ' for c in line):
                if i + 1 < len(lines):
                    next_line = lines[i + 1].strip()
                    if next_line and all(c in '|-: ' for c in next_line):
                        i += 2
                        continue

            cleaned_lines.append(lines[i])
            i += 1

        return '\n'.join(cleaned_lines)

    def _convert_chart_to_markdown(self, chart):
        """Returns (markdown, ok) where ok=False signals that the chart could
        not be converted structurally and the slide should be rendered for
        LLM transcription as a fallback."""
        try:
            md = "\n\n### Chart"
            if chart.has_title:
                md += f": {chart.chart_title.text_frame.text}"
            md += "\n\n"
            data = []
            category_names = [c.label for c in chart.plots[0].categories]
            series_names = [s.name for s in chart.series]
            data.append(["Category"] + series_names)

            for idx, category in enumerate(category_names):
                row = [category]
                for series in chart.series:
                    row.append(series.values[idx])
                data.append(row)

            markdown_table = []
            for row in data:
                markdown_table.append("| " + " | ".join(map(str, row)) + " |")
            header = markdown_table[0]
            separator = "|" + "|".join(["---"] * len(data[0])) + "|"
            return md + "\n".join([header, separator] + markdown_table[1:]), True
        except ValueError as e:
            if "unsupported plot type" in str(e):
                return "\n\n[unsupported chart]\n\n", False
            return "\n\n[unsupported chart]\n\n", False
        except Exception:
            return "\n\n[unsupported chart]\n\n", False

    @staticmethod
    def _transcribe_slides(
        pptx_bytes: bytes,
        slide_indices: list[int],
        *,
        client,
        model,
        prompt,
    ) -> dict[int, str]:
        try:
            from . import _pptx_slide_render
        except ImportError:
            return {}

        with tempfile.NamedTemporaryFile(
            suffix=".pptx", delete=False
        ) as tmp:
            tmp.write(pptx_bytes)
            tmp_path = tmp.name

        try:
            try:
                pngs = _pptx_slide_render.render_slides(tmp_path, slide_indices)
            except Exception:
                return {}

            results: dict[int, str] = {}
            for idx, png_bytes in pngs.items():
                stream_info = StreamInfo(mimetype="image/png", extension=".png")
                try:
                    text = llm_caption(
                        io.BytesIO(png_bytes),
                        stream_info,
                        client=client,
                        model=model,
                        prompt=prompt,
                    )
                except Exception:
                    text = None
                if text:
                    results[idx] = text
            return results
        finally:
            try:
                os.remove(tmp_path)
            except Exception:
                pass
